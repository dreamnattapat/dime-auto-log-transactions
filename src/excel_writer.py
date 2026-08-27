from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

COLUMNS = [
    "received_date",
    "effective_date",
    "settlement_date",
    "issue_date",
    "order_id",
    "transaction_type",
    "security",
    "exchange",
    "units",
    "unit_price",
    "currency",
    "gross_amount",
    "vat",
    "withholding_tax",
    "total_amount",
    "gross_amount_thb",
    "vat_thb",
    "withholding_tax_thb",
    "total_amount_thb",
    "fx_rate_thb_usd",
    "total_fee_thb",
    "account_no",
    "tax_invoice_no",
    "parse_status",
    "raw_text_snippet",
    "gmail_message_id",
]


def append_transactions(excel_path: Path, rows: list[dict]) -> None:
    if excel_path.exists():
        wb = load_workbook(excel_path)
        ws = wb.active
    else:
        excel_path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Transactions"
        ws.append(COLUMNS)

    for row in rows:
        ws.append([row.get(col, "") for col in COLUMNS])
    wb.save(excel_path)


def build_rows(gmail_message_id: str, internal_date_ms: int, transactions: list[dict]) -> list[dict]:
    received_date = datetime.fromtimestamp(internal_date_ms / 1000, tz=timezone.utc).isoformat()
    rows = []
    for fields in transactions:
        row = {col: fields.get(col) for col in COLUMNS}
        row["received_date"] = received_date
        row["gmail_message_id"] = gmail_message_id
        rows.append(row)
    return rows
